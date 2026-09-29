import json,re,csv,datetime
from statistics import median
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
data=json.load(open('merged.json'))
def slug_label(s):
    s=re.sub(r'^(pokemon|one-piece)-','',s or '')
    return ' '.join(w.capitalize() for w in s.split('-'))
MAIN={'Box','Pack','Coffret'}
for m in data:
    if m.get('tcgurl'):
        m['dname']=m['name']; m['dset']=m.get('tset') or ''
    else:
        lab=slug_label(m.get('pcset')); m['dset']=lab
        m['dname']=m['name'].strip() if m['name'].strip().lower().startswith(lab.lower()) else f"{lab} {m['name'].strip()}"
data.sort(key=lambda m:(m['game'],m['cat'],m['dset'].lower(),m['dname'].lower()))
F=Font(name='Arial',size=10); FB=Font(name='Arial',size=10,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78')
MONEY='$#,##0.00'
wb=Workbook()
st=wb.active; st.title='Settings'
st['A1']='Settings'; st['A1'].font=Font(name='Arial',size=12,bold=True)
st['A2']='Discount (Reduced = Median Price x (1 - Discount))'; st['B2']=0.30; st['B2'].number_format='0%'
st['B2'].font=Font(name='Arial',size=10,color='0000FF'); st['B2'].fill=PatternFill('solid',fgColor='FFFF00')
st['A3']='Kraken discount (Kraken Price = Median Price x (1 - Kraken discount))'; st['B3']=0.30; st['B3'].number_format='0%'
st['B3'].font=Font(name='Arial',size=10,color='0000FF'); st['B3'].fill=PatternFill('solid',fgColor='FFFF00')
st['A4']='Edit the yellow cells only. Every "Reduced" and "Kraken Price" value on the other sheets recalculates from them.'
st['A5']='Data collected'; st['B5']=datetime.date.today().isoformat()
st['A6']='Sources'; st['B6']='PriceCharting (ungraded/sealed sale price, USD) - https://www.pricecharting.com ; TCGplayer (Market, Median, Lowest listing, USD) - https://www.tcgplayer.com'
st['A11']='Kraken Price'; st['B11']='Median Price minus the Kraken discount (Settings!B3, 30% by default): the buy/offer price used by the shop.'
st['A7']='Price'; st['B7']='PriceCharting price when available, otherwise TCGplayer Market (then Median, then Lowest).'
st['A8']='Median Price'; st['B8']='MEDIAN of all price points found for the product (PriceCharting price, TCGplayer Market, Median and Lowest). Products found on both sites are merged into one row.'
st['A9']='Categories'; st['B9']='Box = booster box / display. Pack = single boosters and blisters. Coffret = elite trainer boxes, collections, tins, bundles, gift sets. Decks, cases and accessories are on the "Other Sealed" sheet.'
st['A10']='Scope'; st['B10']='Pokemon (English, Japanese, other languages as listed by the sites) and One Piece Card Game sealed products. Products with no price on either site were left out.'
for r in range(1,11):
    for c in 'AB':
        cell=st[f'{c}{r}']
        if cell.font.name!='Arial' or r>4: cell.font=F if not (c=='B' and r in (2,3)) else cell.font
st.column_dimensions['A'].width=48; st.column_dimensions['B'].width=110
for r in range(5,12): st[f'B{r}'].alignment=Alignment(wrap_text=True,vertical='top')
HEAD=['Name','Price','Median Price','Reduced','Kraken Price','Game','Category','Set','PriceCharting Price','TCGplayer Market','TCGplayer Median','TCGplayer Lowest','Sources','PriceCharting Link','TCGplayer Link']
WID=[62,12,13,12,13,10,12,34,17,16,16,16,24,60,50]
def fill(ws,rows):
    ws.append(HEAD)
    for i,h in enumerate(HEAD,1):
        c=ws.cell(row=1,column=i); c.font=FB; c.fill=HF; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=WID[i-1]
    for r,m in enumerate(rows,2):
        ws.cell(r,1,m['dname'])
        ws.cell(r,2,f'=IF(I{r}<>"",I{r},IF(J{r}<>"",J{r},IF(K{r}<>"",K{r},L{r})))')
        ws.cell(r,3,f'=MEDIAN(I{r}:L{r})')
        ws.cell(r,4,f'=ROUND(C{r}*(1-Settings!$B$2),2)')
        ws.cell(r,5,f'=ROUND(C{r}*(1-Settings!$B$3),2)')
        ws.cell(r,6,m['game']); ws.cell(r,7,m['cat']); ws.cell(r,8,m['dset'])
        for col,key in ((9,'pc'),(10,'tm'),(11,'tmed'),(12,'tlow')):
            if m.get(key): ws.cell(r,col,round(float(m[key]),2))
        ws.cell(r,13,' + '.join(sorted(m['sources'])))
        if m.get('pcurl'): c=ws.cell(r,14,m['pcurl']); c.hyperlink=m['pcurl']; c.font=Font(name='Arial',size=10,color='0563C1',underline='single')
        if m.get('tcgurl'): c=ws.cell(r,15,m['tcgurl']); c.hyperlink=m['tcgurl']; c.font=Font(name='Arial',size=10,color='0563C1',underline='single')
        for col in range(1,14):
            c=ws.cell(r,col)
            c.font=F
            if col in (2,3,4,5,9,10,11,12): c.number_format=MONEY
    ws.freeze_panes='B2'; ws.auto_filter.ref=f"A1:{get_column_letter(len(HEAD))}{len(rows)+1}"
main=[m for m in data if m['cat'] in MAIN]; other=[m for m in data if m['cat'] not in MAIN]
fill(wb.create_sheet('Prices',0),main); fill(wb.create_sheet('Other Sealed',1),other)
wb.save('pokemon-onepiece-sealed-prices.xlsx')
# CSV with computed values (same logic as formulas)
with open('pokemon-onepiece-sealed-prices.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(HEAD[:13])
    for m in main:
        pts=[float(x) for x in (m.get('pc'),m.get('tm'),m.get('tmed'),m.get('tlow')) if x]
        price=next(float(x) for x in (m.get('pc'),m.get('tm'),m.get('tmed'),m.get('tlow')) if x)
        med=median(pts)
        w.writerow([m['dname'],f'{price:.2f}',f'{med:.2f}',f'{med*0.7:.2f}',f'{med*0.7:.2f}',m['game'],m['cat'],m['dset']]+[f'{float(m[k]):.2f}' if m.get(k) else '' for k in ('pc','tm','tmed','tlow')]+[' + '.join(sorted(m['sources']))])
import collections
print('main',len(main),collections.Counter((m['game'],m['cat']) for m in main)); print('other',len(other))
